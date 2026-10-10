import json
from io import BytesIO
from dataclasses import dataclass
from urllib.error import URLError

import pytest
from docx import Document

from app import create_app
from app.controllers import tool_controller
from app.services import seo_checker_service
from app.services.seo_checker_service import PageSeoParser
from database import get_seo_check_history_result, record_seo_check_history
from database.common import get_connection


def test_page_seo_parser_collects_core_seo_fields():
    parser = PageSeoParser()
    parser.feed(
        """
        <html>
          <head>
            <title>Example Title</title>
            <meta name="description" content="Example meta description">
            <meta property="og:title" content="Open Graph Title">
            <meta name="twitter:card" content="summary">
            <link rel="canonical" href="https://example.com/page">
          </head>
          <body>
            <h1>Main Heading</h1>
            <h2><span>Section</span> Heading</h2>
            <h3>First Detail</h3>
            <h3>Second Detail</h3>
            <h2>Next Section</h2>
            <img src="/image.jpg">
            <img src="/logo.jpg" alt="Example logo">
            <a href="/internal">Internal</a>
          </body>
        </html>
        """
    )

    assert parser.title == "Example Title"
    assert parser.meta_description == "Example meta description"
    assert parser.canonical == "https://example.com/page"
    assert parser.open_graph["og:title"] == "Open Graph Title"
    assert parser.twitter_cards["twitter:card"] == "summary"
    assert len(parser.headings["h1"]) == 1
    assert parser.headings["h2"] == ["Section Heading", "Next Section"]
    assert parser.heading_sequence == [
        {"level": "h1", "text": "Main Heading"},
        {"level": "h2", "text": "Section Heading"},
        {"level": "h3", "text": "First Detail"},
        {"level": "h3", "text": "Second Detail"},
        {"level": "h2", "text": "Next Section"},
    ]
    assert len(parser.images) == 2
    assert not parser.images[0]["has_alt"]
    assert parser.images[1]["has_alt"]
    assert parser.images[1]["alt"] == "Example logo"
    assert parser.links[0]["href"] == "/internal"
    assert parser.links[0]["text"] == "Internal"
    assert parser.links[0]["region"] == "content"


def test_page_seo_parser_marks_header_and_footer_links():
    parser = PageSeoParser()
    parser.feed(
        """
        <html>
          <body>
            <header><a id="site-logo" href="/header-link"><img id="logo-image" src="/logo.png" alt="Example logo"></a></header>
            <main><a href="/content-link">Content Link</a></main>
            <footer><a href="/footer-link">Footer Link</a></footer>
          </body>
        </html>
        """
    )

    links_by_href = {link["href"]: link for link in parser.links}

    assert links_by_href["/header-link"]["region"] == "header"
    assert links_by_href["/header-link"]["id"] == "site-logo"
    assert links_by_href["/header-link"]["is_image"] is True
    assert links_by_href["/header-link"]["image_alts"] == ["Example logo"]
    assert links_by_href["/header-link"]["image_ids"] == ["logo-image"]
    assert links_by_href["/content-link"]["region"] == "content"
    assert links_by_href["/footer-link"]["region"] == "footer"


def test_page_seo_parser_collects_json_ld_schema():
    parser = PageSeoParser()
    parser.feed(
        """
        <html>
          <head>
            <script type="application/ld+json">
              {"@context":"https://schema.org","@type":"Article","headline":"Example"}
            </script>
          </head>
          <body><h1>Example</h1></body>
        </html>
        """
    )

    assert len(parser.schema_scripts) == 1
    assert '"@type":"Article"' in parser.schema_scripts[0]


def test_social_previews_use_platform_metadata_and_absolute_images():
    parser = PageSeoParser()
    parser.feed(
        """
        <html>
          <head>
            <title>Google result title</title>
            <meta name="description" content="Google result description">
            <meta property="og:title" content="Facebook title">
            <meta property="og:description" content="Facebook description">
            <meta property="og:image" content="/images/social.jpg">
            <meta property="og:site_name" content="Example Site">
            <meta name="twitter:card" content="summary_large_image">
            <meta name="twitter:title" content="Twitter title">
          </head>
        </html>
        """
    )

    previews = seo_checker_service._build_social_preview(parser, "https://example.com/page")

    assert previews["google"]["title"] == "Google result title"
    assert previews["facebook"]["title"] == "Facebook title"
    assert previews["facebook"]["image"] == "https://example.com/images/social.jpg"
    assert previews["twitter"]["title"] == "Twitter title"
    assert previews["twitter"]["image"] == "https://example.com/images/social.jpg"


def test_schema_reports_keep_all_blocks_and_complete_content():
    schemas = [
        json.dumps({"@context": "https://schema.org", "@type": "Article", "description": "x" * 700, "index": index})
        for index in range(25)
    ]

    reports = seo_checker_service._schema_reports(schemas)

    assert len(reports) == 25
    assert reports[-1]["index"] == 25
    assert len(reports[0]["preview"]) < len(reports[0]["content"])
    assert '"description": "' + ("x" * 700) + '"' in reports[0]["content"]


def test_schema_reports_build_nested_hierarchy():
    reports = seo_checker_service._schema_reports(
        [
            json.dumps(
                {
                    "@graph": [
                        {
                            "@type": "WebSite",
                            "@id": "https://example.com/#website",
                            "alternateName": ["Example", "Example Site"],
                        },
                        {
                            "@type": "WebPage",
                            "isPartOf": {"@id": "https://example.com/#website"},
                            "about": [{"@type": "Thing", "name": "SEO"}],
                        },
                    ]
                }
            )
        ]
    )

    hierarchy = reports[0]["hierarchy"]

    assert [row["value"] for row in hierarchy if row["key"] == "@type"] == ["WebSite", "WebPage", "Thing"]
    assert {row["key"] for row in hierarchy if row["depth"] == 1} >= {"0", "1", "@id", "@type", "name"}
    assert next(row for row in hierarchy if row["key"] == "@id")["is_url"] is True


def test_link_report_tracks_404_links(monkeypatch):
    def fake_check_link_url(url, link_type, verify_ssl=True, allow_private=False):
        status_code = 404 if url.endswith("/missing") else 200
        status = "broken" if status_code == 404 else "ok"
        return {"url": url, "type": link_type, "status_code": status_code, "status": status}

    monkeypatch.setattr(seo_checker_service, "_check_link_url", fake_check_link_url)

    report = seo_checker_service._build_link_report(
        [
            {"href": "/missing", "text": "Missing page"},
            {"href": "/working", "text": "Working page"},
        ],
        page_url="https://example.com/page",
        base_url="https://example.com",
    )

    assert report["checked_count"] == 2
    assert report["broken_sample_count"] == 1
    assert report["not_found_count"] == 1
    assert report["broken_count"] == 0
    assert report["not_found_links"][0]["url"] == "https://example.com/missing"
    assert report["not_found_links"][0]["text"] == "Missing page"


def test_link_report_tracks_redirecting_links(monkeypatch):
    def fake_check_link_url(url, link_type, verify_ssl=True, allow_private=False):
        if url.endswith("/old"):
            return {
                "url": url,
                "final_url": "https://example.com/new",
                "type": link_type,
                "status_code": 200,
                "status": "ok",
                "redirected": True,
            }
        return {
            "url": url,
            "final_url": url,
            "type": link_type,
            "status_code": 200,
            "status": "ok",
            "redirected": False,
        }

    monkeypatch.setattr(seo_checker_service, "_check_link_url", fake_check_link_url)

    report = seo_checker_service._build_link_report(
        [
            {"href": "/old", "text": "Old page"},
            {"href": "/current", "text": "Current page"},
        ],
        page_url="https://example.com/page",
        base_url="https://example.com",
    )

    assert report["checked_count"] == 2
    assert report["redirect_count"] == 1
    assert report["redirect_links"][0]["url"] == "https://example.com/old"
    assert report["redirect_links"][0]["final_url"] == "https://example.com/new"


def test_link_report_excludes_header_and_footer_links_from_health_checks(monkeypatch):
    checked_urls = []

    def fake_check_link_url(url, link_type, verify_ssl=True, allow_private=False):
        checked_urls.append(url)
        return {
            "url": url,
            "final_url": url,
            "type": link_type,
            "status_code": 200,
            "status": "ok",
            "redirected": False,
        }

    monkeypatch.setattr(seo_checker_service, "_check_link_url", fake_check_link_url)

    report = seo_checker_service._build_link_report(
        [
            {
                "href": "/header",
                "id": "main-navigation",
                "text": "",
                "region": "header",
                "is_image": True,
                "image_alts": ["Site logo"],
                "image_ids": ["site-logo-image"],
            },
            {"href": "/content", "text": "Content", "region": "content"},
            {"href": "/footer", "text": "Footer", "region": "footer"},
        ],
        page_url="https://example.com/page",
        base_url="https://example.com",
    )

    assert checked_urls == ["https://example.com/content"]
    assert report["checked_count"] == 1
    assert report["header_footer_count"] == 2
    assert [link["region"] for link in report["header_footer_links"]] == ["header", "footer"]
    assert report["header_footer_links"][0]["is_image"] is True
    assert report["header_footer_links"][0]["image_alt"] == "Site logo"
    assert report["header_footer_links"][0]["id"] == "main-navigation"
    assert report["header_footer_links"][0]["image_id"] == "site-logo-image"


def test_header_footer_details_deduplicate_anchor_and_url_pairs():
    repeated_link = {
        "url": "https://example.com/about",
        "text": "About",
        "id": "about-link",
        "type": "internal",
        "region": "header",
        "is_image": False,
        "image_alt": "",
    }
    page_results = [
        {
            "status": "checked",
            "url": f"https://example.com/page-{index}",
            "result": {"stats": {"links": {"header_footer_links": [dict(repeated_link)]}}},
        }
        for index in range(3)
    ]

    details = tool_controller._site_header_footer_link_details(page_results)

    assert details["total_count"] == 1
    assert details["duplicate_count"] == 2
    assert details["rows"][0]["element_id"] == "about-link"


def test_local_urls_are_allowed_only_when_enabled():
    with pytest.raises(ValueError, match="Local/private"):
        seo_checker_service._validate_public_http_url("http://localhost:3000")

    assert seo_checker_service._validate_public_http_url("http://localhost:3000", allow_private=True) is None


@dataclass
class FakeDiscoveryResult:
    base_url: str
    pages: list
    errors: list
    sitemaps: list


@pytest.mark.parametrize(
    ("url", "expected"),
    [
        ("https://example.com/", "Homepage"),
        ("https://example.com/payments/", "/payments"),
        ("https://example.com/slots?sort=new", "/slots"),
        ("https://example.com/blogs/blog-name", "/blogs/blog-name"),
        ("https://example.com/blogs/seo%20guide", "/blogs/seo guide"),
    ],
)
def test_seo_page_tab_label_uses_url_slug(url, expected):
    assert tool_controller._seo_page_tab_label(url) == expected


def test_seo_checker_lists_pages_then_checks_each_page(monkeypatch):
    calls = []

    monkeypatch.setattr(
        tool_controller,
        "discover_website_pages",
        lambda url, limit, allow_private=False: FakeDiscoveryResult(
            base_url="https://example.com",
            pages=["https://example.com/a", "https://example.com/b"],
            errors=[],
            sitemaps=[],
        ),
    )

    def fake_run_seo_audit(url, verify_ssl=True, allow_private=False, use_ai=True):
        calls.append(url)
        return {
            "url": url,
            "status_code": 200,
            "content_type": "text/html",
            "score": 80,
            "grade": "B",
            "checks": [
                {
                    "name": "Meta title",
                    "status": "pass",
                    "detail": "Good title",
                    "recommendation": "Keep it clear.",
                }
            ],
            "ai_summary": {"summary": "Looks usable.", "priority_actions": [], "source": "Rules"},
            "stats": {
                "title": f"Title for {url}",
                "word_count": 500,
                "h1_count": 1,
                "missing_alt_count": 0,
                "links": {
                    "checked_count": 1,
                    "not_found_count": 1 if url.endswith("/a") else 0,
                    "not_found_links": [
                        {
                            "url": f"{url}/missing",
                            "final_url": f"{url}/missing",
                            "type": "internal",
                            "status_code": 404,
                            "status": "broken",
                            "redirected": False,
                            "text": "Missing link",
                        }
                    ] if url.endswith("/a") else [],
                    "broken_count": 0,
                    "broken_links": [],
                    "unreachable_count": 0,
                    "unreachable_links": [],
                    "redirect_count": 1,
                    "redirect_links": [
                        {
                            "url": f"{url}/old",
                            "final_url": f"{url}/new",
                            "type": "internal",
                            "status_code": 200,
                            "status": "ok",
                            "redirected": True,
                            "text": "Old link",
                        }
                    ],
                    "sample_checks": [],
                },
            },
        }

    monkeypatch.setattr(tool_controller, "run_seo_audit", fake_run_seo_audit)

    app = create_app()
    app.testing = True
    response = app.test_client().post(
        "/seo-checker",
        data={"url": "https://example.com", "limit": "2"},
    )

    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert calls == ["https://example.com/a", "https://example.com/b"]
    assert "Listed Pages" in html
    assert "2 redirecting links" in html
    assert "Sample Link Health Details" in html
    assert "https://example.com/a/missing" in html
    assert "Missing link" in html
    assert "Found On Page" in html
    assert "https://example.com/a/old" in html
    assert "https://example.com/a/new" in html
    assert "Checked Pages" in html
    assert "Page Checks" in html
    assert html.index("Listed Pages") < html.index("Page Checks")
    assert "https://example.com/a" in html
    assert "https://example.com/b" in html


def test_site_seo_checks_publish_current_page_progress(monkeypatch):
    progress_messages = []

    monkeypatch.setattr(
        tool_controller,
        "discover_website_pages",
        lambda url, limit, allow_private=False: FakeDiscoveryResult(
            base_url="https://example.com",
            pages=["https://example.com/a", "https://example.com/b"],
            errors=[],
            sitemaps=[],
        ),
    )

    def fake_run_seo_audit(url, verify_ssl=True, allow_private=False, use_ai=True):
        return {
            "url": url,
            "score": 90,
            "grade": "A",
            "checks": [],
            "ai_summary": {"summary": "", "priority_actions": [], "source": "Rules"},
            "stats": {"title": "", "word_count": 0, "h1_count": 0, "missing_alt_count": 0},
        }

    monkeypatch.setattr(tool_controller, "run_seo_audit", fake_run_seo_audit)

    result = tool_controller._run_site_seo_checks(
        "https://example.com",
        limit=2,
        progress_callback=progress_messages.append,
    )

    assert "Checking page 1/2: https://example.com/a" in progress_messages
    assert "Checking page 2/2: https://example.com/b" in progress_messages
    assert result["checked_pages"] == [
        {"index": 1, "url": "https://example.com/a", "status": "checked", "score": 90, "grade": "A", "error": ""},
        {"index": 2, "url": "https://example.com/b", "status": "checked", "score": 90, "grade": "A", "error": ""},
    ]


def test_seo_checker_passes_allow_private_urls(monkeypatch):
    flags = {"discovery": None, "audit": []}

    def fake_discover_website_pages(url, limit, allow_private=False):
        flags["discovery"] = allow_private
        return FakeDiscoveryResult(
            base_url="http://localhost:3000",
            pages=["http://localhost:3000/"],
            errors=[],
            sitemaps=[],
        )

    def fake_run_seo_audit(url, verify_ssl=True, allow_private=False, use_ai=True):
        flags["audit"].append(allow_private)
        return {
            "url": url,
            "score": 90,
            "grade": "A",
            "checks": [],
            "ai_summary": {"summary": "", "priority_actions": [], "source": "Rules"},
            "stats": {"title": "", "word_count": 0, "h1_count": 0, "missing_alt_count": 0, "links": {}},
        }

    monkeypatch.setattr(tool_controller, "discover_website_pages", fake_discover_website_pages)
    monkeypatch.setattr(tool_controller, "run_seo_audit", fake_run_seo_audit)

    app = create_app()
    app.testing = True
    response = app.test_client().post(
        "/seo-checker",
        data={"url": "http://localhost:3000", "limit": "1", "allow_private_urls": "1"},
    )

    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert flags["discovery"] is True
    assert flags["audit"] == [True]
    assert "Allow local/private URLs" in html


def test_seo_checker_v2_renders_internal_links_and_schema(monkeypatch):
    monkeypatch.setattr(
        tool_controller,
        "discover_website_pages",
        lambda url, limit, allow_private=False: FakeDiscoveryResult(
            base_url="https://example.com",
            pages=["https://example.com/a"],
            errors=[],
            sitemaps=[],
        ),
    )

    def fake_run_seo_audit(url, verify_ssl=True, allow_private=False, use_ai=True):
        return {
            "url": url,
            "score": 88,
            "grade": "B",
            "checks": [
                {"name": "Meta title", "status": "pass", "detail": "48 characters", "recommendation": "Keep it clear."},
                {"name": "Meta description", "status": "warn", "detail": "90 characters", "recommendation": "Expand it."},
                {"name": "Image alt text", "status": "pass", "detail": "0 missing", "recommendation": "Maintain alt text."},
                {"name": "Content depth", "status": "pass", "detail": "500 words found", "recommendation": "Keep useful content."},
                {"name": "Heading structure", "status": "pass", "detail": "1 H1, 2 H2", "recommendation": "Keep structure."},
                {"name": "Canonical URL", "status": "pass", "detail": "https://example.com/a", "recommendation": "Keep canonical."},
                {"name": "Sitemap", "status": "pass", "detail": "Sitemap found", "recommendation": "Keep sitemap."},
                {"name": "Robots.txt", "status": "pass", "detail": "Robots.txt found", "recommendation": "Keep robots.txt."},
                {"name": "Social cards", "status": "pass", "detail": "2 Open Graph tags", "recommendation": "Keep social cards."},
                {"name": "Internal links", "status": "pass", "detail": "1 internal", "recommendation": "Keep links useful."},
                {"name": "Sample link health", "status": "warn", "detail": "1 redirect", "recommendation": "Review redirects."},
            ],
            "ai_summary": {"summary": "", "priority_actions": [], "source": "Rules"},
            "stats": {
                "title": "Page A",
                "title_length": 6,
                "meta_description": "A useful page description.",
                "meta_description_length": 26,
                "word_count": 500,
                "h1_count": 1,
                "h2_count": 2,
                "heading_sequence": [
                    {"level": "h1", "text": "Page A"},
                    {"level": "h2", "text": "SEO Details"},
                    {"level": "h3", "text": "Nested Topic"},
                ],
                "canonical": "https://example.com/a",
                "robots_url": "https://example.com/robots.txt",
                "robots_status_code": 200,
                "sitemaps": [{"url": "https://example.com/sitemap.xml", "found": True, "status_code": 200}],
                "open_graph": {"og:title": "Page A", "og:description": "Social description"},
                "twitter_cards": {"twitter:card": "summary"},
                "missing_alt_count": 0,
                "links": {
                    "internal_count": 1,
                    "external_count": 1,
                    "checked_count": 3,
                    "not_found_count": 1,
                    "unreachable_count": 0,
                    "redirect_count": 1,
                    "not_found_links": [
                        {
                            "url": "https://example.com/missing-page",
                            "final_url": "https://example.com/missing-page",
                            "type": "internal",
                            "status_code": 404,
                            "status": "broken",
                            "redirected": False,
                            "text": "Missing page",
                            "id": "missing-page-link",
                        }
                    ],
                    "header_footer_links": [
                        {
                            "url": "https://example.com/",
                            "text": "",
                            "id": "site-logo-link",
                            "type": "internal",
                            "region": "header",
                            "is_image": True,
                            "image_alt": "Example site logo",
                            "image_id": "site-logo-image",
                        }
                    ],
                    "sample_checks": [
                        {
                            "url": "https://example.com/old",
                            "final_url": "https://example.com/new",
                            "type": "internal",
                            "status_code": 200,
                            "status": "ok",
                            "redirected": True,
                            "text": "Old internal page",
                        },
                        {
                            "url": "https://external.example/",
                            "final_url": "https://external.example/",
                            "type": "external",
                            "status_code": 200,
                            "status": "ok",
                            "redirected": False,
                            "text": "External",
                        },
                    ],
                },
                "schemas": [
                    {
                        "index": 1,
                        "valid": True,
                        "type": "Article",
                        "error": "",
                        "preview": '{"@type":"Article"}',
                        "content": '{\n  "@type": "Article",\n  "headline": "Complete schema content"\n}',
                        "hierarchy": [
                            {"key": "@type", "value": "Article", "depth": 0, "is_url": False},
                            {"key": "@id", "value": "https://example.com/#article", "depth": 0, "is_url": True},
                            {"key": "headline", "value": "Complete schema content", "depth": 0, "is_url": False},
                        ],
                    }
                ],
            },
        }

    monkeypatch.setattr(tool_controller, "run_seo_audit", fake_run_seo_audit)

    app = create_app()
    app.testing = True
    client = app.test_client()
    response = client.post(
        "/seo-checker-v2",
        data={"url": "https://example.com", "limit": "1"},
    )
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert "Website SEO Checker V2" in html
    assert "Links From Header And Footer" in html
    assert "Only links inside semantic header and footer elements are listed here." in html
    assert "Example site logo" in html
    assert "ID: site-logo-link" in html
    assert "Image ID: site-logo-image" in html
    assert "Missing alt text" not in html
    assert "Old internal page" in html
    assert "https://example.com/new" in html
    assert "Schema For Each Page" in html
    assert "Article" in html
    assert "Read more" in html
    assert "Complete schema content" in html
    assert "Headings Hierarchy" in html
    assert "&lt;H1&gt;" in html
    assert "Nested Topic" in html
    assert 'href="https://example.com/#article"' in html
    assert "data-link-pagination" in html
    assert 'aria-label="Link Health pagination"' in html
    assert "Meta title" in html
    assert "Meta description" in html
    assert "Image alt text" in html
    assert "Content depth" in html
    assert "Heading structure" in html
    assert "Canonical URL" in html
    assert "Sitemap" in html
    assert "Robots.txt" in html
    assert "Social cards" in html
    assert "Internal links" in html
    assert "Sample link health" in html
    assert "Social Media View" in html
    assert "Google Search" in html
    assert "Facebook" in html
    assert "Twitter / X" in html
    assert "Detected Metadata" in html
    assert "Open Graph" in html
    assert "Twitter Cards" in html
    assert "Download Complete Report" in html
    assert "Download Performance Summary" in html
    assert "All 404 Links" in html
    assert "https://example.com/missing-page" in html
    assert 'aria-label="All 404 links pagination"' in html
    assert "max-h-80" in html
    assert html.count("data-audit-table") >= 5
    assert "data-table-search" in html
    assert "data-table-filter" in html
    assert "Open full report" in html
    assert "Report Contents" in html
    assert 'href="#seo-v2-navigation"' in html
    assert 'href="#seo-v2-pages"' in html
    assert 'href="#seo-v2-404"' in html
    assert "data-seo-toc-link" in html
    assert "IntersectionObserver" in html

    with get_connection() as connection:
        saved_row = connection.execute(
            "SELECT id FROM seo_check_history WHERE checker_version = 'v2' AND base_url = ? ORDER BY id DESC LIMIT 1",
            ("https://example.com",),
        ).fetchone()
    assert saved_row is not None

    restored = client.get(f"/seo-checker-v2?history_id={saved_row['id']}")
    restored_html = restored.get_data(as_text=True)
    assert restored.status_code == 200
    assert "Loaded complete JSON report saved on" in restored_html
    assert "Complete schema content" in restored_html

    with get_connection() as connection:
        connection.execute("DELETE FROM seo_check_history WHERE id = ?", (saved_row["id"],))


def test_v2_history_round_trips_complete_nested_json():
    result = {
        "mode": "site",
        "source_url": "https://json-history.example",
        "base_url": "https://json-history.example",
        "limit": 25,
        "summary": {"average_score": 91, "average_grade": "A", "checked_count": 1},
        "pages": [
            {
                "url": "https://json-history.example/page",
                "result": {
                    "stats": {
                        "schemas": [{"content": '{"@type":"Article"}', "hierarchy": [{"key": "@type", "value": "Article"}]}],
                        "links": {"sample_checks": [{"url": "https://json-history.example/next", "status_code": 200}]},
                    }
                },
            }
        ],
    }

    history_id = record_seo_check_history(result, checker_version="v2")
    try:
        snapshot = get_seo_check_history_result(history_id, checker_version="v2")
        assert snapshot is not None
        assert snapshot["result"] == result
        assert get_seo_check_history_result(history_id, checker_version="v1") is None
    finally:
        with get_connection() as connection:
            connection.execute("DELETE FROM seo_check_history WHERE id = ?", (history_id,))


def test_seo_checker_v2_ai_suggestions_checkbox_controls_audit(monkeypatch):
    ai_flags = []

    def fake_site_check(
        url,
        limit=10,
        verify_ssl=True,
        allow_private=False,
        use_ai=True,
        progress_callback=None,
    ):
        ai_flags.append(use_ai)
        return {
            "mode": "empty",
            "source_url": url,
            "base_url": url,
            "limit": limit,
            "use_ai": use_ai,
            "summary": {},
        }

    monkeypatch.setattr(tool_controller, "_run_site_seo_checks", fake_site_check)
    monkeypatch.setattr(tool_controller, "record_seo_check_history", lambda result, checker_version="v1": 1)
    monkeypatch.setattr(tool_controller, "get_previous_seo_check_history", lambda *args, **kwargs: None)
    monkeypatch.setattr(tool_controller, "list_seo_check_history", lambda *args, **kwargs: [])

    app = create_app()
    app.testing = True
    client = app.test_client()

    without_ai = client.post("/seo-checker-v2", data={"url": "https://example.com", "limit": "1"})
    with_ai = client.post(
        "/seo-checker-v2",
        data={"url": "https://example.com", "limit": "1", "use_ai": "1"},
    )

    assert without_ai.status_code == 200
    assert with_ai.status_code == 200
    assert ai_flags == [False, True]
    assert 'name="use_ai"' in with_ai.get_data(as_text=True)
    assert 'name="use_ai" value="1" checked' in with_ai.get_data(as_text=True)


def test_seo_checker_auto_allows_localhost_ports(monkeypatch):
    flags = {"discovery": None, "audit": []}

    def fake_discover_website_pages(url, limit, allow_private=False):
        flags["discovery"] = allow_private
        return FakeDiscoveryResult(
            base_url="http://localhost:3001",
            pages=["http://localhost:3001/"],
            errors=[],
            sitemaps=[],
        )

    def fake_run_seo_audit(url, verify_ssl=True, allow_private=False, use_ai=True):
        flags["audit"].append(allow_private)
        return {
            "url": url,
            "score": 90,
            "grade": "A",
            "checks": [],
            "ai_summary": {"summary": "", "priority_actions": [], "source": "Rules"},
            "stats": {"title": "", "word_count": 0, "h1_count": 0, "missing_alt_count": 0, "links": {}},
        }

    monkeypatch.setattr(tool_controller, "discover_website_pages", fake_discover_website_pages)
    monkeypatch.setattr(tool_controller, "run_seo_audit", fake_run_seo_audit)

    app = create_app()
    app.testing = True
    response = app.test_client().post(
        "/seo-checker",
        data={"url": "http://localhost:3001", "limit": "1"},
    )

    assert response.status_code == 200
    assert flags["discovery"] is True
    assert flags["audit"] == [True]


def test_seo_checker_compares_against_previous_check(monkeypatch):
    scores = [70, 84]
    issue_counts = [2, 1]
    domain = "https://seo-history-compare.example"

    with get_connection() as connection:
        connection.execute("DELETE FROM seo_check_history WHERE base_url = ?", (domain,))

    monkeypatch.setattr(
        tool_controller,
        "discover_website_pages",
        lambda url, limit, allow_private=False: FakeDiscoveryResult(
            base_url=domain,
            pages=[f"{domain}/page"],
            errors=[],
            sitemaps=[],
        ),
    )

    def fake_run_seo_audit(url, verify_ssl=True, allow_private=False, use_ai=True):
        score = scores.pop(0)
        issue_count = issue_counts.pop(0)
        checks = [
            {"name": "Meta title", "status": "pass", "detail": "Good", "recommendation": "Keep it."},
        ]
        checks.extend(
            {"name": f"Issue {index}", "status": "warn", "detail": "Needs work", "recommendation": "Fix it."}
            for index in range(issue_count)
        )
        return {
            "url": url,
            "score": score,
            "grade": tool_controller._seo_grade(score),
            "checks": checks,
            "ai_summary": {"summary": "", "priority_actions": [], "source": "Rules"},
            "stats": {
                "title": "History page",
                "word_count": 500,
                "h1_count": 1,
                "missing_alt_count": 0,
                "links": {
                    "checked_count": 1,
                    "not_found_count": 1 if score == 70 else 0,
                    "not_found_links": [],
                    "broken_count": 0,
                    "broken_links": [],
                    "unreachable_count": 0,
                    "unreachable_links": [],
                    "redirect_count": 0,
                    "redirect_links": [],
                    "sample_checks": [],
                },
                "schemas": [],
            },
        }

    monkeypatch.setattr(tool_controller, "run_seo_audit", fake_run_seo_audit)

    app = create_app()
    app.testing = True
    client = app.test_client()

    try:
        first = client.post("/seo-checker", data={"url": domain, "limit": "1"})
        second = client.post("/seo-checker", data={"url": domain, "limit": "1"})
        html = second.get_data(as_text=True)

        assert first.status_code == 200
        assert second.status_code == 200
        assert "Compared With Last Check" in html
        assert "Score improved by 14 points" in html
        assert "+14/100" in html
        assert "Warning/fix items" in html
        assert "-1" in html
        assert "SEO Checker History" in html
    finally:
        with get_connection() as connection:
            connection.execute("DELETE FROM seo_check_history WHERE base_url = ?", (domain,))


def test_fetch_url_retries_localhost_through_docker_host(monkeypatch):
    requested_urls = []

    class FakeHeaders:
        def get(self, key, default=""):
            return "text/html" if key.lower() == "content-type" else default

        def get_content_charset(self):
            return "utf-8"

    class FakeResponse:
        status = 200
        headers = FakeHeaders()

        def __init__(self, url):
            self.url = url

        def __enter__(self):
            return self

        def __exit__(self, exc_type, exc, traceback):
            return False

        def geturl(self):
            return self.url

        def read(self, limit):
            return b"<html><title>Local app</title><body><h1>Local app</h1></body></html>"

    def fake_urlopen(request, timeout=None, context=None):
        requested_urls.append(request.full_url)
        if request.full_url == "http://localhost:3001/":
            raise URLError(ConnectionRefusedError("refused"))
        return FakeResponse(request.full_url)

    monkeypatch.setattr(seo_checker_service, "urlopen", fake_urlopen)

    result = seo_checker_service.fetch_url("http://localhost:3001/", allow_private=True)

    assert requested_urls == ["http://localhost:3001/", "http://host.docker.internal:3001/"]
    assert result.url == "http://localhost:3001/"
    assert "Local app" in result.text


def test_seo_checker_download_report_returns_docx():
    result = {
        "mode": "site",
        "base_url": "https://example.com",
        "summary": {
            "average_grade": "B",
            "average_score": 82,
            "discovered_count": 2,
            "checked_count": 2,
            "issue_count": 1,
            "checked_link_count": 4,
            "not_found_link_count": 1,
            "redirect_link_count": 1,
        },
        "pages": [
            {
                "url": "https://example.com/a",
                "status": "checked",
                "result": {
                    "grade": "B",
                    "score": 82,
                    "stats": {"title": "Page A"},
                    "checks": [
                        {
                            "name": "Sample link health",
                            "status": "warn",
                            "detail": "1 404",
                            "recommendation": "Fix missing link.",
                        }
                    ],
                },
            }
        ],
        "link_health_details": {
            "not_found_links": [
                {
                    "source_page": "https://example.com/a",
                    "url": "https://example.com/missing",
                    "final_url": "https://example.com/missing",
                    "status_code": 404,
                    "text": "Missing link",
                    "type": "internal",
                }
            ],
            "redirect_links": [
                {
                    "source_page": "https://example.com/a",
                    "url": "https://example.com/old",
                    "final_url": "https://example.com/new",
                    "status_code": 200,
                    "text": "Old link",
                    "type": "internal",
                }
            ],
        },
    }

    app = create_app()
    app.testing = True
    response = app.test_client().post("/seo-checker/download-report", data={"result_json": json.dumps(result)})

    assert response.status_code == 200
    assert response.headers["Content-Type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    doc = Document(BytesIO(response.data))
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    table_text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    combined = f"{text}\n{table_text}"
    assert "Website SEO Checker Report" in combined
    assert "Detailed Link Health" in combined
    assert "Complete Audit Documentation" in combined
    assert "Page-By-Page Documentation" in combined
    assert "All 404 Links" in combined
    assert "https://example.com/missing" in combined
    assert "Missing link" in combined


def test_seo_checker_download_performance_summary_returns_html():
    result = {
        "mode": "site",
        "source_url": "https://summary.example",
        "base_url": "https://summary.example",
        "summary": {
            "average_grade": "B",
            "average_score": 84,
            "discovered_count": 1,
            "checked_count": 1,
            "issue_count": 1,
            "not_found_link_count": 1,
            "schema_count": 1,
            "valid_schema_count": 1,
        },
        "pages": [
            {
                "url": "https://summary.example/page",
                "status": "checked",
                "result": {
                    "url": "https://summary.example/page",
                    "status_code": 200,
                    "score": 84,
                    "grade": "B",
                    "checks": [
                        {"name": "Meta title", "status": "pass", "recommendation": "Keep the title clear."},
                        {"name": "Meta description", "status": "warn", "recommendation": "Expand the meta description."},
                    ],
                    "stats": {
                        "title_length": 52,
                        "word_count": 900,
                        "ssl_verified": True,
                        "schemas": [{"valid": True}],
                    },
                },
            }
        ],
    }

    app = create_app()
    app.testing = True
    response = app.test_client().post(
        "/seo-checker/download-performance-summary",
        data={"result_json": json.dumps(result)},
    )
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert response.headers["Content-Type"].startswith("text/html")
    assert "attachment;" in response.headers["Content-Disposition"]
    assert response.headers["Content-Disposition"].endswith(".html")
    assert "SEO Performance Summary - summary.example" in html
    assert "Reported Average Score: 84/100" in html
    assert "Executive Assessment" in html
    assert "Recommended Maintenance Focus" in html
    assert "@page { size: A4 landscape" in html
    assert "slotcasino.site" not in html
